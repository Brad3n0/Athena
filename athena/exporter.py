"""Export chats (Markdown / Word) and back up or restore all of Athena's data."""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import PurePosixPath
from typing import Any

from . import store

SKIP_DIRS = {"browser-profile"}  # the desktop app's browser cache — large and not needed


class ExportError(Exception):
    pass


def safe_name(title: str | None) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "", (title or "Athena chat")).strip()[:80] or "Athena chat"


def _messages(chat: dict[str, Any]):
    for m in chat.get("messages", []):
        if m.get("role") in ("user", "assistant") and (m.get("content") or "").strip():
            text = m.get("display") if m.get("role") == "user" and m.get("display") else m.get("content")
            yield ("You" if m["role"] == "user" else "Athena"), text.strip()


def export_chat(chat: dict[str, Any], fmt: str) -> tuple[bytes, str, str]:
    title = chat.get("title") or "Athena chat"
    when = datetime.fromtimestamp(chat.get("updated") or chat.get("created") or 0).strftime("%B %d, %Y")
    if fmt == "md":
        lines = [f"# {title}", f"_Exported from Athena AI · {when}_", ""]
        for who, text in _messages(chat):
            lines += [f"### {who}", "", text, ""]
        return "\n".join(lines).encode("utf-8"), "text/markdown", "md"
    if fmt == "docx":
        try:
            import docx
            from docx.shared import Pt, RGBColor
        except ImportError:
            raise ExportError("Word export needs 'python-docx' — restart Athena with start.bat to install it")
        doc = docx.Document()
        doc.add_heading(title, level=1)
        sub = doc.add_paragraph(f"Exported from Athena AI · {when}")
        sub.runs[0].italic = True
        for who, text in _messages(chat):
            head = doc.add_paragraph()
            run = head.add_run(who)
            run.bold = True
            run.font.color.rgb = RGBColor(0xB0, 0x7D, 0x12) if who == "Athena" else RGBColor(0x33, 0x33, 0x33)
            in_code = False
            for line in text.splitlines():
                if line.strip().startswith("```"):
                    in_code = not in_code
                    continue
                p = doc.add_paragraph()
                r = p.add_run(line)
                if in_code:
                    r.font.name = "Consolas"
                    r.font.size = Pt(9.5)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    raise ExportError(f"Unknown format '{fmt}'")


def make_backup() -> bytes:
    buf = io.BytesIO()
    base = store.DATA_DIR
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("athena-backup.txt", f"Athena AI backup · {datetime.now():%Y-%m-%d %H:%M}\n")
        if base.exists():
            for path in base.rglob("*"):
                rel = path.relative_to(base)
                if path.is_file() and not (set(rel.parts) & SKIP_DIRS) and not path.name.endswith(".tmp"):
                    z.write(path, rel.as_posix())
    return buf.getvalue()


def restore_backup(data: bytes) -> int:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ExportError("That isn't an Athena backup (.zip) file")
    names = z.namelist()
    if "athena-backup.txt" not in names:
        raise ExportError("That zip isn't an Athena backup")
    count = 0
    for name in names:
        rel = PurePosixPath(name)
        if name.endswith("/") or name == "athena-backup.txt":
            continue
        if rel.is_absolute() or ".." in rel.parts or (set(rel.parts) & SKIP_DIRS):
            continue  # never write outside the data folder
        target = store.DATA_DIR.joinpath(*rel.parts)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(z.read(name))
        count += 1
    return count
